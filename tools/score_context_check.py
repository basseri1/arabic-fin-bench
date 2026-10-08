"""Score the blind check of the ground truth's context (gt_ctx): row labels and periods reconstructed from the page
images by an independent annotator, against gt/.

Each returned workbook has the rows of the ground truth in order (row 10 onwards), with the figures given; the
annotator typed every row label (column B), every column header (row 7) and every period (row 8). The script first
checks that rows and figures are unchanged, then compares:
  labels  of rows with figures, normalised as in the label metric (eval2.norm_ar) and compared by string similarity:
          same (>= 95), close (80-95), different (< 80), missing (annotator left it empty) or extra (the ground
          truth has no label). Section headings are compared too, but reported apart.
  periods of every value column: the end date (or the year, if the annotator gave only a year). Columns of equity
          matrices, where the ground truth records no period, are compared only if the annotator gave one.
Everything that is not 'same' goes to results/adjudication.csv for a decision against the page image: GT_WRONG,
AUDIT_WRONG, BOTH_WRONG or EQUIVALENT (both name the same printed line; the difference is spelling or typing).
Decisions already entered are kept. Once all are decided, the report gives the share of rows with figures whose
ground-truth label is wrong, and of columns whose period is wrong, with exact one-sided 95% upper bounds.

usage: python tools/score_context_check.py [--returned gt_ctx/returned] [--results gt_ctx/results]
"""
import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from rapidfuzz import fuzz

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gtlib  # noqa: E402
from score_audit import cp_upper  # noqa: E402

E = gtlib.eval2
PAPER = gtlib.PAPER
FIRST = 10
DECISIONS = {"GT_WRONG", "AUDIT_WRONG", "BOTH_WRONG", "EQUIVALENT"}
FIELDS = ["filing_id", "sheet", "statement", "pages", "excel_ref", "item", "gt", "annotator", "similarity", "class",
          "figures_in_row", "annotator_comment", "decision", "adjudicator_note"]


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.lower() in ("-", "–", "—", "nil"):                       # the ground truth stores a printed dash as "nil"
        return 0.0
    vals = E.numbers(s, signed=True)
    return vals[0] if vals else None


def period_of(v):
    """('date', 'YYYY-MM-DD') or ('year', 'YYYY') or None, from what the annotator typed (Excel may store a date)."""
    if v is None or v == "":
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return ("date", v.strftime("%Y-%m-%d"))
    s = E.norm_digits(str(v))
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return ("date", f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}")
    m = re.search(r"(19|20)\d\d", s)
    return ("year", m.group(0)) if m else None


def label_class(g, a):
    g, a = E.norm_ar(g or ""), E.norm_ar(a or "")
    if not g and not a:
        return "same", 100
    if not a:
        return "missing", 0
    if not g:
        return "extra", 0
    s = fuzz.ratio(g, a)
    return ("same" if s >= 95 else "close" if s >= 80 else "different"), round(s)


def score_workbook(path):
    wb = load_workbook(path, data_only=True)
    meta = {r[0].value: r[1].value for r in wb["filing"].iter_rows(min_row=2) if r[0].value}
    fid, who = meta.get("filing_id"), (meta.get("annotator") or "").strip()
    gt = json.load(open(PAPER / "gt" / f"{fid}.json", encoding="utf-8"))
    lines, problems = [], []
    for ws, st in zip(wb.worksheets[2:], gt["statements"]):
        pages = ", ".join(map(str, st["pages"]))
        for j, r in enumerate(st["rows"]):
            x = FIRST + j
            if (ws.cell(row=x, column=1).value or "") != r["kind"]:
                problems.append(f"{ws.title} row {x}: row kind changed")
            vals = r.get("values") or {}
            for k, col in enumerate(st["columns"]):
                g = vals.get(col["id"])
                if g is not None and num(ws.cell(row=x, column=4 + k).value) != num(g):
                    problems.append(f"{ws.title} {ws.cell(row=x, column=4 + k).coordinate}: figure changed")
            cls, sim = label_class(r.get("label_ar"), ws.cell(row=x, column=2).value)
            lines.append(dict(filing_id=fid, sheet=ws.title, statement=st["type"], pages=pages, excel_ref=f"B{x}",
                              item="label" if vals else "heading", gt=r.get("label_ar") or "",
                              annotator=ws.cell(row=x, column=2).value or "", similarity=sim, cls=cls,
                              figures_in_row=sum(v is not None for v in vals.values()),
                              annotator_comment=ws.cell(row=x, column=20).value or ""))
        for k, col in enumerate(st["columns"]):
            c = 4 + k
            g = col.get("period")
            a = period_of(ws.cell(row=8, column=c).value)
            if not g and not a:
                continue                                                 # equity component column: not assessed
            if not a:
                cls = "missing"
            elif not g:
                cls = "extra"
            else:
                cls = "same" if (a[1] == g if a[0] == "date" else a[1] == g[:4]) else "different"
            nfig = sum(1 for r in st["rows"] if (r.get("values") or {}).get(col["id"]) is not None)
            lines.append(dict(filing_id=fid, sheet=ws.title, statement=st["type"], pages=pages,
                              excel_ref=ws.cell(row=8, column=c).coordinate, item="period", gt=g or "",
                              annotator=a[1] if a else "", similarity="", cls=cls, figures_in_row=nfig,
                              annotator_comment=f"header typed: {ws.cell(row=7, column=c).value or ''}"))
    return fid, who, lines, problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--returned", default=str(PAPER / "gt_ctx" / "returned"))
    ap.add_argument("--results", default=str(PAPER / "gt_ctx" / "results"))
    a = ap.parse_args()
    returned, results = Path(a.returned), Path(a.results)
    results.mkdir(parents=True, exist_ok=True)
    books = sorted(p for p in returned.glob("*.xlsx") if not p.name.startswith("~$"))
    if not books:
        raise SystemExit(f"no workbooks in {returned}")
    all_lines, problems, who, fids = [], [], set(), []
    for p in books:
        fid, w, lines, probs = score_workbook(p)
        fids.append(fid)
        who.add(w or "(name not given)")
        all_lines += lines
        problems += [f"{fid}: {x}" for x in probs]

    adj = results / "adjudication.csv"
    key = lambda d: (d["filing_id"], d["sheet"], d["excel_ref"], d["item"])
    old = {key(d): d for d in csv.DictReader(open(adj, encoding="utf-8-sig"))} if adj.exists() else {}
    review = [d for d in all_lines if d["cls"] != "same"]
    with open(adj, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for d in review:
            prev = old.get(key(d), {})
            d["decision"], d["adjudicator_note"] = prev.get("decision", ""), prev.get("adjudicator_note", "")
            w.writerow({**{k: d.get(k, "") for k in FIELDS}, "class": d["cls"]})

    def block(item, title):
        L = [d for d in all_lines if d["item"] == item]
        if not L:
            return []
        n = len(L)
        same = sum(d["cls"] == "same" for d in L)
        R = [d for d in L if d["cls"] != "same"]
        dec = [d for d in R if d.get("decision", "").strip().upper() in DECISIONS]
        gt_err = sum(d["decision"].strip().upper() in ("GT_WRONG", "BOTH_WRONG") for d in dec)
        cls = {c: sum(d["cls"] == c for d in R) for c in ("close", "different", "missing", "extra")}
        out = [f"### {title}", "",
               f"Compared: {n}. Same: {same} ({100 * same / n:.1f}%). To adjudicate: {len(R)} "
               f"({', '.join(f'{k} {v}' for k, v in cls.items() if v) or 'none'}).", ""]
        if len(dec) == len(R):
            figs = sum(d["figures_in_row"] for d in dec if d["decision"].strip().upper() in ("GT_WRONG", "BOTH_WRONG"))
            out.append(f"All decided. Ground-truth errors: **{gt_err} of {n}** ({100 * gt_err / n:.2f}%, exact one-sided "
                       f"95% upper bound {100 * cp_upper(gt_err, n):.2f}%), affecting {figs} figures.")
        else:
            out.append(f"{len(R) - len(dec)} of {len(R)} still need a decision in `adjudication.csv`.")
        return out + [""]

    out = ["# Context check of the ground truth: results", "",
           f"Workbooks scored: {len(books)} ({', '.join(sorted(fids))}). Annotator: {', '.join(sorted(who))}.", ""]
    if problems:
        out += ["**Structure changed in the returned workbooks (fix before trusting the scores):**", ""] + \
               [f"- {x}" for x in problems[:50]] + [""]
    out += block("label", "Row labels of rows with figures") + block("period", "Periods of value columns") + \
        block("heading", "Section headings (no figures; reported apart)")
    (results / "CONTEXT_REPORT.md").write_text("\n".join(out) + "\n", encoding="utf-8")
    print("\n".join(out))
    print(f"wrote {results / 'CONTEXT_REPORT.md'} and {adj}")


if __name__ == "__main__":
    main()
