"""Worked examples of the fact scorer's verdicts, for checking the scoring by eye (local use only: it embeds crops of
the filings' pages, which are not redistributed).

For each example: a crop of the printed page around the figure's line, the ground-truth lines around it (the
transcription), the rows the system wrote around the figure (exactly as written), the scorer's verdict with the step
that decided it, and, where one exists, the human checker's or the five test annotators' judgment. Examples cover each
verdict (complete, misplaced, unresolved), cases the scorer gets right and cases it gets wrong.
Writes agent_check/results/scoring_examples.html.   usage: python tools/scoring_examples.py
"""
import base64
import html
import json
import re
import sys
from pathlib import Path

import fitz
from openpyxl import load_workbook

PAPER = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PAPER / "bench"))
sys.path.insert(0, str(PAPER / "tools"))
import common_subset as CS  # noqa: E402
import rq2_strict as R  # noqa: E402

RUNS = {"Cohere Parse": CS.SYSTEMS["Cohere Parse"], "Mistral OCR": CS.SYSTEMS["Mistral OCR"],
        "dots.mocr": R.MAIN["dots.mocr pipeline, seed 0"][0], "Chandra OCR 2": R.MAIN["Chandra OCR 2, own input size"][0]}
AGENT_SYS = {"Cohere Parse": "Cohere Parse", "Mistral OCR": "Mistral OCR", "dots.mocr pipeline, seed 0": "dots.mocr",
             "Chandra OCR 2, own input size": "Chandra OCR 2"}
DIG = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")

# (verdict group, scorer right?, item, title, what to look at)
EXAMPLES = [
    ("Complete", True, "E046", "Label and header both name the printed line and year",
     "Mistral OCR writes the label beside the figures and the year above them; the scorer matches both and credits the figure."),
    ("Unresolved", True, "E002", "Same page, another system: the table has no labels at all",
     "dots.mocr keeps the figures, note numbers and section headings of the Almarai balance sheet but drops every line-item "
     "label, so no figure can be tied to its line. The verdict is right: nothing in the output names the line."),
    ("Complete (new rule)", True, "S122", "A subtotal printed without a label, written without a label",
     "The printed line has no label (an operating subtotal). Cohere Parse writes the figures in a row without a label. Until "
     "today this was 'row not stated'; under the rule adopted today it is complete, and the checker's judgment agrees."),
    ("Misplaced: row", True, "E013", "Labels shifted one line against the figures",
     "Chandra OCR 2 pairs each label with the figures of the line above it. Every figure is right and every total still adds "
     "up, but each figure sits under its neighbour's label. The verdict is right."),
    ("Misplaced: row", True, "E066", "Figures attached to a section heading",
     "Mistral OCR writes a row of figures on the heading 'items that will or may be reclassified...' and moves the labels below "
     "it up by one. The figure sits under the wrong label. The verdict is right."),
    ("Misplaced: row (scorer wrong)", False, "S206", "A short label marked as the wrong line",
     "dots.mocr writes a shortened label for earnings per share. The label is right for a reader, but the scorer's row "
     "alignment matches it to another line, so the figure is counted as wrong. The checker marked it complete."),
    ("Misplaced: period (scorer wrong)", False, "E056", "A year row one cell short",
     "Cohere Parse's year row has one cell fewer than the table, so read cell by cell, each year sits one column off. A reader "
     "assigns the years to the figure columns in order and gets the right year; the scorer reads the shifted header and "
     "counts the figure under the wrong period. All five test annotators judged the period right."),
    ("Unresolved (scorer wrong)", False, "S065", "A partial label that still names the line",
     "Cohere Parse's label leaves out a word ('zakat') of the printed label. It still names the line, but its similarity "
     "falls below the scorer's threshold, so the figure is 'row not stated'. The checker marked it complete."),
    ("Unresolved (scorer wrong)", False, "S200", "Unresolved, but in fact under the wrong period",
     "Chandra OCR 2's header names three dates over two columns of figures. The scorer cannot assign them and leaves the "
     "period unresolved; the checker found that the figure in fact belongs to 31 December 2023, not to the date above it, "
     "so the figure is wrong, not merely unresolved."),
    ("Unresolved: period (debatable)", None, "E001", "The same shifted header: unresolved, though the years in order are right",
     "Same Chandra table. Its header has three dates (31 Dec 2024, 31 Dec 2023, 1 Jan 2023) but only two columns of figures. "
     "Assigning the first two dates in order gives the right year for this figure; the scorer refuses to guess and leaves it "
     "unresolved. Three of the five test annotators judged the period right."),
    ("Attribution (debatable)", None, "E117", "The same value printed in both years",
     "The share capital of 300,000 thousand riyals is printed for both years. The benchmark matches output figures to "
     "ground-truth cells by value, one to one, so the figure written under 2021 is matched to the 2022 cell. The output is "
     "right; the error lies in the matching."),
]


# ------------------------------------------------------------------ locating a ground-truth line on the printed page
def number_tokens(page):
    """Numbers in the text layer with their baseline y: (y, x0, x1, {values})."""
    chars = []
    for b in page.get_text("rawdict")["blocks"]:
        for l in b.get("lines", []):
            for s in l["spans"]:
                for ch in s["chars"]:
                    if ch["c"] in "0123456789٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹,.٫٬،":
                        chars.append((round(ch["origin"][1]), ch["bbox"][0], ch["bbox"][2], ch["c"], s["size"]))
    lines = {}
    for y, x0, x1, c, size in chars:
        key = next((k for k in lines if abs(k - y) <= 2), y)
        lines.setdefault(key, []).append((x0, x1, c, size))
    import textlayer
    toks = []
    for y, chs in lines.items():
        chs.sort()
        tok, start, last = "", None, None
        for x0, x1, c, size in chs + [(1e9, 1e9, " ", 0)]:
            if last is not None and x0 - last > max(1.5, 0.3 * size):
                v = textlayer.parse(tok)
                if v:
                    toks.append((y, start, last, v))
                tok, start = "", None
            tok += c
            start = x0 if start is None else start
            last = x1
    return toks


def locate(page, values):
    """Baseline y of the printed line carrying these values (most of them on one baseline)."""
    toks = number_tokens(page)
    best = None
    for v in values:
        for y, _, _, vals in toks:
            if any(abs(abs(x) - abs(v)) < 0.51 for x in vals):
                score = sum(1 for w in values for y2, _, _, vv in toks if abs(y2 - y) <= 3 and any(abs(abs(x) - abs(w)) < 0.51 for x in vv))
                if best is None or score > best[0]:
                    best = (score, y)
    return best[1] if best else None


def crop(fid, pno, values, box=None):
    doc = fitz.open(PAPER / "filings" / f"{fid}.pdf")
    page = doc[pno - 1]
    W, H = page.rect.width, page.rect.height
    if box is None:
        y = locate(page, values) if page.get_text("words") else None
        if y is None:                                            # scanned page: the whole page, smaller
            pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4))
            return base64.b64encode(pix.tobytes("png")).decode()
        page.draw_rect(fitz.Rect(18, y - 10, W - 18, y + 4), color=(0.85, 0.45, 0), fill=(1, 0.85, 0.2), fill_opacity=0.3, width=0.8)
        rect = fitz.Rect(10, max(0, y - 95), W - 10, min(H, y + 75))
    else:
        rect = fitz.Rect(*box)
    pix = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2), clip=rect)
    return base64.b64encode(pix.tobytes("png")).decode()


# ------------------------------------------------------------------ rendering
def e(x):
    return html.escape(str(x if x is not None else ""))


def fmt(v):
    if v is None:
        return ""
    if not isinstance(v, (int, float)):
        return str(v)                                            # printed dashes, 'nil'
    s = f"{abs(v):,.0f}" if float(v).is_integer() else f"{abs(v):,.2f}"
    return f"({s})" if v < 0 else s


def gt_table(st, rid, cid, also=()):
    rows = [r for r in st["rows"]]
    pos = next(i for i, r in enumerate(rows) if r["id"] == rid)
    lo, hi = max(0, pos - 3), min(len(rows), pos + 4)
    cols = st["columns"]
    h = ['<table class="t" dir="rtl"><tr><th>line (as printed)</th>' + "".join(f"<th>{e(c.get('label', ''))}</th>" for c in cols) + "</tr>"]
    for r in rows[lo:hi]:
        cls = ' class="target"' if r["id"] == rid else ""
        lab = r.get("label_ar", "")
        if r.get("kind") == "section":
            lab = f"<b>{e(lab)}</b>"
        elif not lab:
            lab = '<span class="none">(no label printed)</span>'
        else:
            lab = e(lab)
        cells = "".join(f'<td class="{"fig" if (r["id"] == rid and c["id"] == cid) else ""}">{e(fmt((r.get("values") or {}).get(c["id"])))}</td>'
                        for c in cols)
        h.append(f"<tr{cls}><td>{lab}</td>{cells}</tr>")
    return "".join(h) + "</table>"


def out_table(d, cell):
    ri, j, first = d["raw"][cell["i"]], d["cols"][cell["c"]], d["first"]
    W = max(len(g) for g in d["grid"])
    h = ['<table class="t" dir="rtl">']
    for r in range(max(0, first - 3), first):
        h.append('<tr class="hdr">' + "".join(f"<td>{e(' '.join(str(x or '').split()))}</td>" for x in d["grid"][r]) + "</tr>")
    if first == 0:
        h.append(f'<tr class="hdr"><td colspan="{W}"><span class="none">(no header: the table starts with figures)</span></td></tr>')
    lo, hi = max(first, ri - 3), min(len(d["grid"]), ri + 4)
    if lo > first:
        h.append(f'<tr><td colspan="{W}" class="none">… {lo - first} rows</td></tr>')
    for r in range(lo, hi):
        tds = []
        for jj, x in enumerate(d["grid"][r]):
            x = " ".join(str(x or "").split())
            cls = "fig" if (r == ri and jj == j) else ""
            tds.append(f'<td class="{cls}">{e(x) if x else ""}</td>')
        h.append(f'<tr class="{"target" if r == ri else ""}">' + "".join(tds) + "</tr>")
    if hi < len(d["grid"]):
        h.append(f'<tr><td colspan="{W}" class="none">… {len(d["grid"]) - hi} rows</td></tr>')
    return "".join(h) + "</table>"


def explain(cell, st_all, k):
    """The scorer's decision for this figure, step by step."""
    steps = []
    claimed = [tuple(x) for x in (cell.get("claimed") or [])]
    lab = cell["label"].strip()
    steps.append(f"Value: {'right' if cell.get('value_ok') else 'wrong'} ({e(fmt(cell['signed']))}).")
    if claimed:
        si, cid = next(((s, c) for s, c in claimed if s == k[0]), claimed[0])
        col = next(c for c in st_all[si]["columns"] if c["id"] == cid)
        steps.append(f"Period: the column header «{e(cell['header']) or '—'}» is read as «{e(col.get('label', ''))}».")
    else:
        steps.append(f"Period: no period can be read from the column header «{e(cell['header']) or '—'}».")
    rs = cell.get("row_status")
    rk = tuple(cell.get("row_key") or ())
    if rs == "unlabeled line":
        steps.append("Line: the output row has no label and the printed line has none either, so the line is stated (rule adopted 3 Oct 2026).")
    elif rk:
        line = next(r for r in st_all[rk[0]]["rows"] if r["id"] == rk[1])
        steps.append(f"Line: the row label «{e(lab)}» is matched to the printed line «{e(line.get('label_ar', ''))}».")
    elif rs == "no label":
        steps.append("Line: the output row has no label.")
    else:
        steps.append(f"Line: the row label «{e(lab)}» matches no printed line closely enough (similarity below the threshold).")
    return steps


def main():
    S = json.load(open(PAPER / "bench" / "rq2_strict_summary.json"))
    tr, tc = S["thresholds"]["row"], S["thresholds"]["col"]
    second = R.MAIN["dots.mocr pipeline, seed 0"][0]
    test = R.score.load_gt(str(PAPER / "gt"), "test")
    akey = json.load(open(PAPER / "agent_check" / "key" / "key.json"))
    skey = json.load(open(PAPER / "scorer_check" / "key" / "key.json"))
    resc = json.load(open(PAPER / "scorer_check" / "key" / "rescored.json"))
    wb = load_workbook(PAPER / "scorer_check" / "results" / "scorer_check.xlsx", data_only=True)
    human = {r[0]: r for r in wb["check"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    agents = {}
    for f in sorted((PAPER / "agent_check" / "returned").glob("annotator_*.json")):
        for x in json.load(open(f))["items"]:
            agents.setdefault(x["item"], []).append((f.stem.split("_")[1], x))
    cache, assigned = {}, {}
    sections = []
    for n, (group, right, item, title, text) in enumerate(EXAMPLES, 1):
        if item.startswith("E"):
            a = akey[item]
            sysn, fid, k = AGENT_SYS[a["system"]], a["filing"], tuple(a["k"])
        else:
            a = skey[item]
            sysn, fid = a["system"], a["filing"]
        run = RUNS[sysn]
        res_key = (run, fid)
        if res_key not in assigned:
            res = R.analyse(run, second, second, None, {fid: test[fid]}, tr, tc, cache)
            recs = []
            CS.outcomes(res, {fid: test[fid]}, recs)
            assigned[res_key] = (res, recs)
        res, recs = assigned[res_key]
        if item.startswith("E"):
            rec = next(r for r in recs if r["k"] == k)
            cell = rec["cell"]
            outcome = "complete" if rec["outcome"] == "complete" else f"{rec['outcome']}: {'period' if rec['side'] == 'column' else rec['side']}"
        else:
            same = [c for c in res["cells"] if c["page"] == a["page"] and abs(c["value"] - abs(a["value"])) <= R.TOL
                    and c["label"] == a["label"] and c["header"] == a["header"]]
            cell = next((c for c in same if c["fact"] == resc[item]["fact"]), same[0])
            if cell.get("ctx"):
                k = tuple(cell["ctx"])
            else:
                cl = [tuple(x) for x in (cell.get("claimed") or [])] or [(si, c["id"]) for si, st in enumerate(test[fid]["statements"])
                                                                         if R.page_no(a["page"]) in st["pages"] for c in st["columns"]]
                _, gtc = R.gt_index(test[fid])
                k = next(kk for kk, v in gtc.items() if v is not None and abs(v - cell["signed"]) <= R.TOL and (kk[0], kk[2]) in cl)
            outcome = {"correct": "complete", "duplicate": "complete"}.get(cell["fact"], cell["fact"])
        gt = test[fid]
        st = gt["statements"][k[0]]
        d = cache[(run, fid)][cell["doc"]]
        pno = R.page_no(cell["page"])
        line = next(r for r in st["rows"] if r["id"] == k[1])
        img = crop(fid, pno, [v for v in (line.get("values") or {}).values() if isinstance(v, (int, float)) and v],
                   BOXES.get(item))
        judge = ""
        if item in human:
            r = human[item]
            judge = (f"<p class='judge'><b>Human checker</b> (scorer check, by hand): printed line «{e(r[9])}», period «{e(r[10])}»; "
                     f"label names it: <b>{e(r[11])}</b>; header names the period: <b>{e(r[12])}</b>; figure right: <b>{e(r[13])}</b>."
                     + (f" Note: “{e(r[14])}”" if r[14] else "") + "</p>")
        if item in agents:
            votes = [f"{x['row']} / {x['period']}" for _, x in agents[item]]
            from collections import Counter
            judge += ("<p class='judge'><b>Five test annotators</b> (AI agents, row / period): " +
                      "; ".join(f"{v} ×{c}" for v, c in Counter(votes).most_common()) + "</p>")
        badge = {True: '<span class="ok">scorer right</span>', False: '<span class="bad">scorer wrong</span>',
                 None: '<span class="mid">debatable</span>'}[right]
        steps = "".join(f"<li>{s}</li>" for s in explain(cell, gt["statements"], k))
        sections.append(f"""
<section>
  <h2><span class="num">{n}</span> {e(title)}</h2>
  <p class="meta">{e(group)} · {badge} · {e(sysn)} · {e(gt.get('entity', fid))}, {e(st.get('title_ar', ''))}, page {pno}</p>
  <p>{e(text)}</p>
  {'<img src="data:image/png;base64,' + img + '" alt="printed page">' if img else '<p class="none">(page crop not available)</p>'}
  <div class="grid">
    <div><h3>Printed (your ground truth)</h3>{gt_table(st, k[1], k[2])}</div>
    <div><h3>What the system wrote</h3>{out_table(d, cell)}</div>
  </div>
  <div class="verdict"><b>Scorer: {e(outcome)}</b><ol>{steps}</ol></div>
  {judge}
</section>""")
    # ---------------------------------------------------------- part B: wrong figures drawn at random
    import random
    rng = random.Random(20261003)
    randoms = []
    for sysn, nn in (("Chandra OCR 2", 4), ("Mistral OCR", 4), ("dots.mocr", 2), ("Cohere Parse", 2)):
        run = RUNS[sysn]
        pool = []
        for fid in sorted(test):
            if (run, fid) not in assigned:
                res = R.analyse(run, second, second, None, {fid: test[fid]}, tr, tc, cache)
                recs = []
                CS.outcomes(res, {fid: test[fid]}, recs)
                assigned[(run, fid)] = (res, recs)
            doc = fitz.open(PAPER / "filings" / f"{fid}.pdf")
            for rec in assigned[(run, fid)][1]:
                if rec["outcome"] == "misplaced" and doc[R.page_no(rec["cell"]["page"]) - 1].get_text("words"):
                    pool.append((fid, rec))
        rng.shuffle(pool)
        seen = set()
        for fid, rec in pool:
            if len([x for x in randoms if x[0] == sysn]) == nn:
                break
            if (fid, rec["cell"]["page"]) in seen:
                continue
            seen.add((fid, rec["cell"]["page"]))
            randoms.append((sysn, run, fid, rec))
    rsec = []
    for m, (sysn, run, fid, rec) in enumerate(randoms, 1):
        k, cell = rec["k"], rec["cell"]
        gt = test[fid]
        st = gt["statements"][k[0]]
        d = cache[(run, fid)][cell["doc"]]
        pno = R.page_no(cell["page"])
        line = next(r for r in st["rows"] if r["id"] == k[1])
        img = crop(fid, pno, [v for v in (line.get("values") or {}).values() if isinstance(v, (int, float)) and v])
        side = "period" if rec["side"] == "column" else rec["side"]
        steps = "".join(f"<li>{x}</li>" for x in explain(cell, gt["statements"], k))
        rsec.append(f"""
<section>
  <h2><span class="num">R{m}</span> {e(sysn)}: counted as misplaced ({e(side)})</h2>
  <p class="meta">{e(gt.get('entity', fid))}, {e(st.get('title_ar', ''))}, page {pno} · drawn at random · you judge</p>
  {'<img src="data:image/png;base64,' + img + '" alt="printed page">' if img else ''}
  <div class="grid">
    <div><h3>Printed (your ground truth)</h3>{gt_table(st, k[1], k[2])}</div>
    <div><h3>What the system wrote</h3>{out_table(d, cell)}</div>
  </div>
  <div class="verdict"><b>Scorer: misplaced ({e(side)})</b><ol>{steps}</ol></div>
</section>""")
    import score_scorer_check as SC
    keyc = json.load(open(PAPER / "scorer_check" / "key" / "key.json"))
    pairs = []
    for it, r in human.items():
        if it not in keyc:
            continue
        fact, grp, unl = resc[it]["fact"], resc[it]["group"], resc[it]["unlabeled_line"]
        eq = fact == "correct, column not assessed"
        h = SC.human_class(r[11], r[12], r[13], eq, unl)
        if h is not None and not eq:
            pairs.append((grp, h))
    rowsum = ""
    for g in ("complete", "wrong context", "unresolved"):
        tot = [h for s_, h in pairs if s_ == g]
        bad = sum(1 for h in tot if h != g)
        rowsum += f"<tr><td>{g if g != 'wrong context' else 'misplaced (wrong line or period)'}</td><td>{len(tot)}</td><td>{bad} ({100 * bad / len(tot):.1f}%)</td><td>" + \
                  ", ".join(f"{c}: {sum(1 for h in tot if h == c)}" for c in ("complete", "wrong context", "unresolved", "wrong value") if c != g and any(h == c for h in tot)) + "</td></tr>"
    how = (f"<section><h2>How often the scorer is wrong (human check by hand, {len(pairs)} figures, current rule)</h2>"
           "<table class='sum'><tr><th>Scorer verdict</th><th>Figures checked</th><th>Checker disagrees</th><th>Checker's verdict instead</th></tr>"
           + rowsum + "</table><p class='meta'>From scorer_check/results/scorer_check.xlsx, re-scored with today's rule for lines printed without a label.</p></section>")
    part_b = ("<h1 style='margin-top:40px'>Part B: figures the scorer counts as wrong, drawn at random</h1>"
              "<p class='intro'>Twelve misplaced figures drawn at random from the common base (Chandra OCR 2 and Mistral OCR four each, "
              "dots.mocr and Cohere Parse two each; one per page; pages with a text layer only, so that the printed line can be found "
              "and highlighted). No commentary: judge each one by eye.</p>" + "".join(rsec))
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Scoring examples</title>
<style>
:root {{ --bg:#fbfaf7; --fg:#1d1d1b; --muted:#6b6a65; --line:#dedbd2; --hl:#fff1b8; --fig:#ffd75e; --ok:#1f7a3a; --bad:#b3261e; --mid:#8a6d00; --card:#ffffff; }}
body {{ background:var(--bg); color:var(--fg); font:15px/1.5 -apple-system, "Segoe UI", sans-serif; margin:0; padding:24px 16px 64px; }}
main {{ max-width:1100px; margin:0 auto; }}
h1 {{ font-size:24px; margin:0 0 6px; }} h2 {{ font-size:18px; margin:0 0 4px; }} h3 {{ font-size:13px; color:var(--muted); margin:0 0 6px; font-weight:600; }}
.intro {{ color:var(--muted); max-width:820px; }}
section {{ background:var(--card); border:1px solid var(--line); border-radius:10px; padding:18px; margin:18px 0; }}
.num {{ display:inline-block; width:26px; height:26px; border-radius:13px; background:var(--fg); color:#fff; text-align:center; font-size:14px; line-height:26px; margin-right:6px; }}
.meta {{ color:var(--muted); font-size:13px; margin:0 0 8px; }}
.ok {{ color:var(--ok); font-weight:600; }} .bad {{ color:var(--bad); font-weight:600; }} .mid {{ color:var(--mid); font-weight:600; }}
img {{ max-width:100%; border:1px solid var(--line); border-radius:6px; margin:6px 0 10px; }}
.grid {{ display:grid; grid-template-columns:1fr 1fr; gap:14px; }}
@media (max-width:820px) {{ .grid {{ grid-template-columns:1fr; }} }}
.t {{ border-collapse:collapse; width:100%; font-family:"Geeza Pro","Noto Naskh Arabic","Arial",sans-serif; font-size:13px; }}
.t td, .t th {{ border:1px solid var(--line); padding:3px 6px; vertical-align:top; }}
.t th {{ background:#f1efe8; font-weight:600; }}
.t tr.hdr td {{ background:#f1efe8; color:var(--muted); }}
.t tr.target td {{ background:var(--hl); }}
.t td.fig {{ background:var(--fig); font-weight:700; outline:2px solid #c99700; }}
.none {{ color:var(--muted); font-style:italic; }}
.verdict {{ margin-top:10px; padding:10px 12px; background:#f6f4ee; border-radius:8px; }}
.verdict ol {{ margin:6px 0 0 18px; padding:0; }}
.judge {{ font-size:14px; margin:8px 0 0; }}
table.sum {{ border-collapse:collapse; margin:10px 0; font-size:14px; }} table.sum td, table.sum th {{ border:1px solid var(--line); padding:4px 8px; text-align:right; }}
table.sum td:first-child, table.sum th:first-child {{ text-align:left; }}
</style></head><body><main>
<h1>How the scorer judges a figure: worked examples</h1>
<p class="intro">Each example shows the printed line (crop of the filing, highlighted), your ground truth around it, the rows
the system wrote around the figure (yellow cell), and the scorer's verdict with the steps that decided it. A figure is
<b>complete</b> when its value is right, its row label names the printed line and its column header names the period;
<b>misplaced</b> when the label or header names another line or period; <b>unresolved</b> when the output names no line or
period that can be matched. Since today, a line printed without a label counts as named when the output row has no label
either. Examples 1–5 show the scorer working as intended; 6–9 show its known errors; 10–11 are debatable.</p>
{how}
{''.join(sections)}
{part_b}
</main></body></html>"""
    out = PAPER / "agent_check" / "results" / "scoring_examples.html"
    out.write_text(page, encoding="utf-8")
    print("wrote", out, f"{out.stat().st_size / 1e6:.1f} MB")


BOXES = {}       # manual crop boxes (PDF points) for pages without a text layer: item -> (x0, y0, x1, y1)


if __name__ == "__main__":
    main()
