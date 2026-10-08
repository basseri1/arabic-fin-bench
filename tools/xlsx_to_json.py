"""Convert an annotated workbook to the JSON ground truth and validate it.

usage: python xlsx_to_json.py paper/annotation/FILING_ID.xlsx [--out paper/gt/FILING_ID.json] [--tol 1]
Exit status 1 when validation finds errors (nothing is written unless --force).
"""
import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gtlib  # noqa: E402
from make_workbook import COMMENT_COL, FIRST, LAST, SUM_COL, VCOLS  # noqa: E402


def text(v):
    return "" if v is None else str(v).strip()


def pages(v):
    out = []
    for part in re.split(r"[,،\s]+", text(v)):
        if re.fullmatch(r"\d+-\d+", part):
            a, b = map(int, part.split("-"))
            out += list(range(a, b + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def sum_refs(v, excel_to_id):
    refs = []
    for part in re.split(r"[,،]", text(v).replace(" ", "")):
        if not part:
            continue
        neg = part.startswith("-")
        part = part.lstrip("-")
        if ":" in part:
            a, b = map(int, part.split(":"))
            rows = range(a, b + 1)
        else:
            rows = [int(part)]
        for x in rows:
            if x in excel_to_id:                     # empty spacer rows inside a range are skipped
                refs.append(("-" if neg else "") + excel_to_id[x])
    return refs


def convert(path):
    wb = load_workbook(path, data_only=False)
    meta = {text(r[0].value): r[1].value for r in wb["filing"].iter_rows(min_row=2) if r[0].value}
    gt = {"schema_version": "1.0", "filing_id": text(meta.get("filing_id")), "entity": text(meta.get("entity")),
          "source": {"url": text(meta.get("source_url")), "sha256": text(meta.get("sha256")),
                     "file": f"filings/{text(meta.get('filing_id'))}.pdf"}}
    for k in ("period_end", "report", "digits", "annotator", "verified_by", "comment"):
        if text(meta.get(k)):
            gt[k] = text(meta.get(k))
    cons = text(meta.get("consolidated")).lower()
    gt["consolidated"] = True if cons == "yes" else (False if cons == "no" else None)
    problems, statements = [], []
    for ws in wb.worksheets:
        if ws.title in ("README", "filing"):
            continue
        data_rows = [x for x in range(FIRST, LAST + 1)
                     if any(text(ws[f"{c}{x}"].value) for c in ["A", "B", "C"] + VCOLS)]
        if not data_rows:
            continue                                   # unused statement sheet
        ccy = text(ws["B4"].value) or "SAR"
        scale = int(float(text(ws["B5"].value) or 1))
        used = [c for c in VCOLS if text(ws[f"{c}7"].value) or any(text(ws[f"{c}{x}"].value) for x in data_rows)]
        cols = []
        for i, c in enumerate(used, 1):
            cd = {"id": f"c{i}", "label": text(ws[f"{c}7"].value), "currency": text(ws[f"{c}9"].value) or ccy}
            if text(ws[f"{c}8"].value):
                cd["period"] = text(ws[f"{c}8"].value)[:10]
            cols.append(cd)
        excel_to_id = {x: f"r{i}" for i, x in enumerate(data_rows, 1)}
        rows = []
        for x in data_rows:
            values = {}
            for c, cd in zip(used, cols):
                raw = ws[f"{c}{x}"].value
                try:
                    v = gtlib.parse_value(raw)
                except ValueError as e:
                    problems.append(f"{ws.title}!{c}{x}: {e}")
                    continue
                if v is not None:
                    values[cd["id"]] = v
            kind = text(ws[f"A{x}"].value) or ("section" if not values else "item")
            row = {"id": excel_to_id[x], "kind": kind, "label_ar": text(ws[f"B{x}"].value)}
            notes = [n.strip() for n in re.split(r"[،,]", text(ws[f"C{x}"].value)) if n.strip()]
            if notes:
                row["notes"] = notes
            if values:
                row["values"] = values
            refs = sum_refs(ws[f"{SUM_COL}{x}"].value, excel_to_id)
            if refs:
                row["sum_of"] = refs
            if text(ws[f"{COMMENT_COL}{x}"].value):
                row["comment"] = text(ws[f"{COMMENT_COL}{x}"].value)
            rows.append(row)
        statements.append({"id": f"s{len(statements) + 1}", "type": text(ws["B1"].value), "title_ar": text(ws["B2"].value),
                           "pages": pages(ws["B3"].value), "unit": {"currency": ccy, "scale": scale},
                           "columns": cols, "rows": rows})
    gt["statements"] = statements
    return gt, problems


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out")
    ap.add_argument("--tol", type=float, default=1.0)
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()
    gt, problems = convert(a.xlsx)
    errors, warnings = gtlib.validate(gt, a.tol)
    errors = problems + errors
    for e in errors:
        print("ERROR  ", e)
    for w in warnings:
        print("warning", w)
    print(f"{gt['filing_id']}: {gtlib.stats(gt)} | {len(errors)} errors, {len(warnings)} warnings")
    if errors and not a.force:
        sys.exit(1)
    out = Path(a.out or gtlib.PAPER / "gt" / f"{gt['filing_id']}.json")
    gtlib.save(gt, out)
    print("wrote", out)
