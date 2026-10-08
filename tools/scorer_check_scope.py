"""Scope-matched reading of the human check of the fact scorer (scorer_check).

1. Complete-fact validation and row-only equity cases apart. The scorer leaves the columns of equity matrices
   unassessed ("correct, column not assessed"); in the main tables these figures count as unresolved. The 8 such items
   in the sample were drawn with the unresolved verdicts but judged on the row only, which moves them to the complete
   group in SCORER_CHECK.md. Here they are reported separately, and agreement and kappa are given without them.
2. Correction with matching scope. The common base (Table 10) holds statements whose columns are periods; equity
   matrices are excluded. The per-verdict rates used to correct its complete-fact shares are recomputed here from the
   checked items of the same scope only (no changes-in-equity statements), and applied to bench/common_subset_summary.json.
Writes scorer_check/results/SCOPE.md.   usage: python tools/scorer_check_scope.py scorer_check/results/scorer_check.xlsx
"""
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from score_scorer_check import CLASSES, cp, header_answer, human_class, kappa, rescored_verdicts  # noqa: E402

PAPER = Path(__file__).resolve().parents[1]
GROUPS = CLASSES[:3]                                   # scorer verdicts: complete, wrong context, unresolved


def main():
    key = json.load(open(PAPER / "scorer_check" / "key" / "key.json"))
    wb = load_workbook(sys.argv[1], data_only=True)
    items = []
    rescored = rescored_verdicts()                         # current scorer verdicts (tools/rescore_scorer_check.py)
    for r in wb["check"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or r[0] not in key:
            continue
        k = key[r[0]]
        fact, group, unl = rescored.get(r[0], (k["scorer"]["fact"], k["group"], False))
        equity_row_only = fact == "correct, column not assessed"
        h = human_class(r[11], header_answer(r[0], r[12], r[10]), r[13], equity_row_only, unl)
        if h is not None:
            items.append(dict(item=r[0], group=group, stype=k["stype"], row_only=equity_row_only, human=h))

    def block(sel, title):
        pairs = [(i["group"], i["human"]) for i in sel]
        n = len(pairs)
        agree = sum(a == b for a, b in pairs)
        conf = Counter(pairs)
        out = [f"## {title}", "", f"Items: {n}. Agreement: {agree} of {n} ({100 * agree / n:.1f}%); Cohen's kappa {kappa(pairs):.3f}.", "",
               "| Scorer \\ Human | " + " | ".join(CLASSES) + " | n | scorer error (95% CI) |", "|---|" + "---:|" * (len(CLASSES) + 2)]
        rates = {}
        for s in GROUPS:
            tot = sum(conf[(s, h)] for h in CLASSES)
            if not tot:
                continue
            e = tot - conf[(s, s)]
            lo, hi = cp(e, tot)
            rates[s] = {h: conf[(s, h)] / tot for h in CLASSES}
            out.append(f"| {s} | " + " | ".join(str(conf[(s, h)]) for h in CLASSES) +
                       f" | {tot} | {e} ({100 * e / tot:.1f}%, {100 * lo:.1f}–{100 * hi:.1f}%) |")
        return out + [""], rates

    def binary(sel, title):
        """Correct (complete fact) against wrong (every other outcome), on both sides, as facts are scored."""
        pairs = [(i["group"] == "complete", i["human"] == "complete") for i in sel]
        n = len(pairs); agree = sum(a == b for a, b in pairs)
        sc = [h for s_, h in pairs if s_]; sw = [h for s_, h in pairs if not s_]
        e1, e2 = sum(1 for h in sc if not h), sum(1 for h in sw if h)
        lo1, hi1 = cp(e1, len(sc)); lo2, hi2 = cp(e2, len(sw))
        return ["## " + title, "", f"Items: {n}. Agreement: {agree} of {n} ({100 * agree / n:.1f}%); Cohen's kappa "
                f"{kappa([(str(a), str(b)) for a, b in pairs]):.3f}.", "",
                f"- Scorer says correct: {len(sc)}; the checker finds {e1} wrong ({100 * e1 / len(sc):.1f}%, 95% CI {100 * lo1:.1f}–{100 * hi1:.1f}%).",
                f"- Scorer says wrong: {len(sw)}; the checker finds {e2} correct ({100 * e2 / len(sw):.1f}%, 95% CI {100 * lo2:.1f}–{100 * hi2:.1f}%).", ""]

    full = [i for i in items if not i["row_only"]]
    row_only = [i for i in items if i["row_only"]]
    same_scope = [i for i in full if i["stype"] != "changes in equity"]
    out = ["# Human check of the fact scorer, by scope", "", "Written by `tools/scorer_check_scope.py`.", ""]
    b, _ = block(full, "1. Complete-fact validation (row-only equity cases excluded)")
    out += b
    ok = sum(i["human"] == "complete" for i in row_only)
    out += ["## 2. Row-only equity cases (column not assessed by the scorer)", "",
            f"Items: {len(row_only)}. The human judged the row right (and the figure right) in {ok} of {len(row_only)}. These "
            "figures count as unresolved in the fact tables; this check validates only their row.", ""]
    b, rates = block(same_scope, "3. Same scope as the common base (statements whose columns are periods; no equity matrices)")
    out += b
    out += binary(full, "3b. Correct against wrong (complete-fact validation, row-only equity cases excluded)")
    out += binary(same_scope, "3c. Correct against wrong, same scope as the common base")
    cs = json.load(open(PAPER / "bench" / "common_subset_summary.json"))["systems"]
    pc = {s: rates[s]["complete"] for s in GROUPS}
    out += ["## 4. Common base (Table 10) corrected with the same-scope rates", "",
            "corrected complete = complete x P(human complete | scorer complete) + unresolved x P(... | unresolved) + "
            f"misplaced x P(... | wrong context) = complete x {pc['complete']:.3f} + unresolved x {pc['unresolved']:.3f} + "
            f"misplaced x {pc['wrong context']:.3f}.", "", "| System | Complete | Corrected | Change (points) |", "|---|---:|---:|---:|"]
    changes = []
    for name, v in cs.items():
        sh = v["shares"]
        corr = sh["complete"] * pc["complete"] + sh["unresolved"] * pc["unresolved"] + sh["misplaced"] * pc["wrong context"]
        changes.append(100 * (corr - sh["complete"]))
        out.append(f"| {name} | {100 * sh['complete']:.1f} | {100 * corr:.1f} | {100 * (corr - sh['complete']):+.1f} |")
    out += ["", f"Largest absolute change: {max(abs(c) for c in changes):.1f} points."]
    res = PAPER / "scorer_check" / "results" / "SCOPE.md"
    res.write_text("\n".join(out) + "\n", encoding="utf-8")
    print("\n".join(out))


if __name__ == "__main__":
    main()
