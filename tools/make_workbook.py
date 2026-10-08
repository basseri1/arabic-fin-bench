"""Create the Excel annotation workbook for one filing (optionally pre-populated from a draft JSON).

Layout of each statement sheet (right-to-left):
  B1 type · B2 title · B3 pages · B4 currency · B5 scale
  row 7  column headers   (D..Q = up to 14 value columns, as printed)
  row 8  period per value column (YYYY-MM-DD)
  row 9  currency per value column (blank = statement currency)
  row 10+ one table row per line: A kind · B label · C notes · D..Q values · R sum_of · S live check · T comment
sum_of lists the Excel rows a total adds up, e.g. "12:18" or "15,27,-30" (a leading minus subtracts). The live
check column shows ✓/✗; validate_gt.py runs the same check on the converted JSON.

usage: python make_workbook.py FILING_ID [--draft path.json] [--out path.xlsx]
"""
import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gtlib  # noqa: E402

FIRST, LAST = 10, 209                           # data rows
VCOLS = [get_column_letter(i) for i in range(4, 18)]   # D..Q
SUM_COL, CHECK_COL, COMMENT_COL = "R", "S", "T"
TYPES = ["financial_position", "income", "comprehensive_income", "income_and_comprehensive_income",
         "changes_in_equity", "changes_in_net_assets", "cash_flows", "other"]
DEFAULT_SHEETS = {
    "Company": ["financial_position", "income", "comprehensive_income", "changes_in_equity", "cash_flows"],
    "Fund": ["financial_position", "comprehensive_income", "changes_in_net_assets", "cash_flows"],
}
PURPLE, PREFILL = "464298", "FFF4CC"
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=PURPLE)
META_FILL = PatternFill("solid", fgColor="EDEBF7")
PREFILL_FILL = PatternFill("solid", fgColor=PREFILL)


def check_formula(r, ncols=len(VCOLS)):
    """Live arithmetic check for row r: each used value column must equal the sum of the rows listed in sum_of
    ("12,15,-18": a leading minus subtracts; "12:18": a range). Only the first ncols value columns are checked."""
    s = f'SUBSTITUTE(${SUM_COL}{r}," ","")'
    parts = []
    for c in VCOLS[:max(1, ncols)]:
        rng = f"{c}${FIRST}:{c}${LAST}"
        listed = (f'SUMPRODUCT(--ISNUMBER(SEARCH(","&ROW({rng})&",",","&{s}&",")),{rng})'
                  f'-SUMPRODUCT(--ISNUMBER(SEARCH(",-"&ROW({rng})&",",","&{s}&",")),{rng})')
        exp = (f'IF(ISNUMBER(SEARCH(":",{s})),'
               f'SUM(INDIRECT("{c}"&LEFT({s},SEARCH(":",{s})-1)&":{c}"&MID({s},SEARCH(":",{s})+1,10))),'
               f'{listed})')
        # a printed dash counts as zero; empty cells are not checked (IF short-circuits, OR would not)
        parts.append(f'IF(OR(ISNUMBER({c}{r}),{c}{r}="-"),ABS(N({c}{r})-({exp}))<=1,TRUE)')
    return f'=IF(${SUM_COL}{r}="","",IF(AND({",".join(parts)}),"✓","✗"))'


def inventory(fid):
    p = gtlib.PAPER / "dataset_inventory.csv"
    for r in csv.DictReader(open(p, encoding="utf-8-sig")):
        if r["filing_id"] == fid:
            return r
    raise SystemExit(f"{fid} is not in dataset_inventory.csv")


def filing_sheet(wb, inv, draft):
    ws = wb.create_sheet("filing")
    ws.sheet_view.rightToLeft = True
    period = (draft or {}).get("period_end", "")
    fields = [
        ("filing_id", inv["filing_id"], "do not change"),
        ("entity", inv["entity"], ""),
        ("source_url", inv.get("public_source_url", ""), ""),
        ("sha256", inv.get("sha256", ""), "fingerprint of the exact PDF annotated"),
        ("period_end", period, "YYYY-MM-DD, e.g. 2024-12-31"),
        ("report", (draft or {}).get("report", "annual" if inv.get("report", "").startswith("Annual") else "interim"), "annual / interim"),
        ("consolidated", {"Yes": "yes", "No": "no"}.get(inv.get("consolidated", ""), ""), "yes / no"),
        ("digits", (draft or {}).get("digits", ""), "arabic-indic / western / mixed — as printed"),
        ("annotator", (draft or {}).get("annotator", ""), "who drafted the figures (the annotator, or your name)"),
        ("verified_by", (draft or {}).get("verified_by", ""), "person who checked every figure against the page image"),
        ("comment", (draft or {}).get("comment", ""), ""),
    ]
    ws.append(["field", "value", "help"])
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for f in fields:
        ws.append(list(f))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 16, 70, 44


def statement_sheet(wb, idx, st, prefilled):
    ws = wb.create_sheet(f"{idx}_{st['type']}"[:31])
    ws.sheet_view.rightToLeft = True
    meta = [("type", st["type"]), ("title_ar", st.get("title_ar", "")),
            ("pages", ",".join(str(p) for p in st.get("pages", []))),
            ("currency", st.get("unit", {}).get("currency", "SAR")), ("scale", st.get("unit", {}).get("scale", 1))]
    for i, (k, v) in enumerate(meta, 1):
        ws[f"A{i}"], ws[f"B{i}"] = k, v
        ws[f"A{i}"].font = Font(bold=True)
        ws[f"B{i}"].fill = META_FILL
    ws["C3"] = "PDF page numbers, e.g. 11 or 11,12"
    ws["C5"] = "1 = riyals, 1000 = thousands, 1000000 = millions (as stated on the page)"
    dv_type = DataValidation(type="list", formula1='"' + ",".join(TYPES) + '"', allow_blank=False)
    dv_scale = DataValidation(type="list", formula1='"1,1000,1000000"', allow_blank=False)
    dv_kind = DataValidation(type="list", formula1='"section,item,subtotal,total"', allow_blank=True)
    for dv, ref in ((dv_type, "B1"), (dv_scale, "B5"), (dv_kind, f"A{FIRST}:A{LAST}")):
        ws.add_data_validation(dv)
        dv.add(ref)
    heads = {"A": "kind", "B": "label_ar", "C": "notes", SUM_COL: "sum_of (rows)", CHECK_COL: "check", COMMENT_COL: "comment"}
    for col, h in heads.items():
        ws[f"{col}7"] = h
    cols = st.get("columns", [])
    for j, c in enumerate(VCOLS):
        cd = cols[j] if j < len(cols) else {}
        ws[f"{c}7"] = cd.get("label", "")
        ws[f"{c}8"] = cd.get("period", "")
        ws[f"{c}9"] = cd.get("currency", "") if cd.get("currency") != st.get("unit", {}).get("currency") else ""
    ws["C8"], ws["C9"] = "period →", "currency →"
    for c in ["A", "B", "C"] + VCOLS + [SUM_COL, CHECK_COL, COMMENT_COL]:
        ws[f"{c}7"].font, ws[f"{c}7"].fill = HEAD_FONT, HEAD_FILL
        ws[f"{c}7"].alignment = Alignment(wrap_text=True, vertical="center")
    for c in VCOLS:
        ws[f"{c}8"].fill = ws[f"{c}9"].fill = META_FILL
    rowpos = {}
    for i, r in enumerate(st.get("rows", [])):
        x = FIRST + i
        rowpos[r["id"]] = x
        ws[f"A{x}"], ws[f"B{x}"] = r["kind"], r.get("label_ar", "")
        ws[f"C{x}"] = "، ".join(r.get("notes", []))
        for j, cd in enumerate(cols[:len(VCOLS)]):
            v = (r.get("values") or {}).get(cd["id"])
            if v is not None:
                ws[f"{VCOLS[j]}{x}"] = "-" if v == "nil" else v
                if prefilled:
                    ws[f"{VCOLS[j]}{x}"].fill = PREFILL_FILL
        if r.get("comment"):
            ws[f"{COMMENT_COL}{x}"] = r["comment"]
        if prefilled:
            ws[f"B{x}"].fill = PREFILL_FILL
    for i, r in enumerate(st.get("rows", [])):
        if r.get("sum_of"):
            ws[f"{SUM_COL}{FIRST + i}"] = ",".join(("-" if s.startswith("-") else "") + str(rowpos[s.lstrip("-")]) for s in r["sum_of"])
    ncols = max(len(cols), 2)
    for x in range(FIRST, LAST + 1):
        ws[f"{CHECK_COL}{x}"] = check_formula(x, ncols)
    red, green = PatternFill("solid", fgColor="F4C7C3"), PatternFill("solid", fgColor="CDEBD3")
    rng = f"{CHECK_COL}{FIRST}:{CHECK_COL}{LAST}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"✗"'], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"✓"'], fill=green))
    widths = {"A": 10, "B": 52, "C": 9, SUM_COL: 14, CHECK_COL: 7, COMMENT_COL: 30}
    for c in VCOLS:
        widths[c] = 16
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    ws.freeze_panes = f"D{FIRST}"


def readme_sheet(wb):
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.rightToLeft = True
    lines = [
        "Annotation workbook — one sheet per financial statement. Full rules: paper/ANNOTATION_GUIDELINE.md",
        "مصنف التوصيف — ورقة لكل قائمة مالية. القواعد الكاملة في دليل التوصيف.",
        "",
        "1. Fill the 'filing' sheet (period_end, digits, your name).",
        "2. On each statement sheet: type, title, PDF pages, currency and scale at the top.",
        "3. Row 7: each value column's header as printed; row 8: its period; row 9: its currency if it differs.",
        "4. One Excel row per printed table row, in printed order. kind: section / item / subtotal / total.",
        "5. Values exactly as printed: negatives in parentheses become minus; a printed dash is typed as -; empty stays empty.",
        "6. Totals: in 'sum_of' list the Excel rows it adds up (12:18 or 15,27). The check column turns green when it adds up.",
        "7. Yellow cells are an unverified annotator draft read from the page images. Verify every one against the page image,",
        "   digit by digit and sign by sign. Most drafting slips were Arabic-Indic ٥ read as ٠ (or the reverse).",
        "8. When the whole filing is checked, put your name in 'verified_by' on the filing sheet.",
        "9. Unused statement sheets: leave them empty; they are skipped.",
    ]
    for l in lines:
        ws.append([l])
    ws.column_dimensions["A"].width = 120
    ws["A1"].font = Font(bold=True, size=12)


def build(fid, draft=None, out=None, prefilled=False):
    inv = inventory(fid)
    wb = Workbook()
    readme_sheet(wb)
    filing_sheet(wb, inv, draft)
    if draft:
        sts = draft["statements"]
    else:
        kinds = DEFAULT_SHEETS["Fund" if inv.get("entity_type") == "Fund" else "Company"]
        sts = [{"type": t, "columns": [{"id": "c1"}, {"id": "c2"}], "rows": [], "unit": {"currency": "SAR", "scale": 1}} for t in kinds]
    for i, st in enumerate(sts, 1):
        statement_sheet(wb, i, st, prefilled)
    out = Path(out or gtlib.PAPER / "annotation" / f"{fid}.xlsx")
    wb.save(out)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("filing_id")
    ap.add_argument("--draft")
    ap.add_argument("--out")
    ap.add_argument("--prefilled", action="store_true", help="mark draft values as pre-filled (yellow)")
    a = ap.parse_args()
    print(build(a.filing_id, gtlib.load(a.draft) if a.draft else None, a.out, a.prefilled))
