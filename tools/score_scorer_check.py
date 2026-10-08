"""Compare a returned human check of the fact scorer with the scorer's verdicts (scorer_check/key/key.json).

Human class, from the checker's answers:
  figure wrong (column N = no)                                   -> wrong value
  row label or header names another line or period (L/M = no)    -> wrong context
  no row label or no readable header (L = no label, M = no header) -> unresolved
  otherwise (L = yes and M = yes)                                -> complete
Scorer group: complete / wrong context / unresolved (see tools/make_scorer_packet.py).
Equity-matrix figures that the scorer leaves "column not assessed" are compared on the row only.
Scorer verdicts are the current ones when scorer_check/key/rescored.json exists (tools/rescore_scorer_check.py). A
figure on a line printed without a label is stated by an output row without a label, on both sides: for such an item
the checker's "no label" counts as the right row. Where the parser fix changed the column header the checker was shown
(Cohere Parse headers with rowspan cells), the header question is answered from the checker's own reading of the
printed period (column K) against the header as now parsed: the same year counts as "yes".

Reports agreement, Cohen's kappa, the confusion matrix, and the scorer's error rate within each verdict group with an
exact 95% interval. Writes scorer_check/results/SCORER_CHECK.md.
usage: python tools/score_scorer_check.py scorer_check/returned/scorer_check_<name>.xlsx
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from scipy.stats import beta

PAPER = Path(__file__).resolve().parents[1]
CLASSES = ("complete", "wrong context", "unresolved", "wrong value")


def cp(k, n, a=0.05):
    lo = 0.0 if k == 0 else beta.ppf(a / 2, k, n - k + 1)
    hi = 1.0 if k == n else beta.ppf(1 - a / 2, k + 1, n - k)
    return lo, hi


def human_class(l, m, nn, equity_row_only, unlabeled_line=False):
    l, m, nn = (str(x or "").strip().lower() for x in (l, m, nn))
    if unlabeled_line and l == "no label":
        l = "yes"                                              # no label printed and none written: the row is right
    if not (l and nn and (m or equity_row_only)):
        return None                                            # unanswered
    if nn == "no":
        return "wrong value"
    if l == "no" or (not equity_row_only and m == "no"):
        return "wrong context"
    if l == "no label" or (not equity_row_only and m == "no header"):
        return "unresolved"
    return "complete"


def kappa(pairs):
    n = len(pairs)
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def rescored_verdicts():
    """item -> (current fact, group, on a line printed without a label), when the check has been re-scored."""
    f = PAPER / "scorer_check" / "key" / "rescored.json"
    if not f.exists():
        return {}
    return {it: (v["fact"], v["group"], v["unlabeled_line"]) for it, v in json.load(open(f)).items()}


YEAR = re.compile(r"(?<![\d٠-٩])(?:19|20|١٩|٢٠)[\d٠-٩]{2}(?![\d٠-٩])")
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def header_answer(item, m, k_period):
    """The checker's header answer, re-read against the header as now parsed when the parser fix changed it."""
    f = PAPER / "scorer_check" / "key" / "rescored.json"
    v = json.load(open(f)).get(item, {}) if f.exists() else {}
    if not v.get("header_changed"):
        return m
    now = {y.translate(DIGITS) for y in YEAR.findall(v.get("header_now") or "")}
    said = {y.translate(DIGITS) for y in YEAR.findall(str(k_period or ""))}
    if not now:
        return "no header"
    return "yes" if said and said <= now and len(now) == 1 else "no"


def main():
    path = Path(sys.argv[1])
    key = json.load(open(PAPER / "scorer_check" / "key" / "key.json"))
    wb = load_workbook(path, data_only=True)
    meta = {str(r[0].value): r[1].value for r in wb["checker"].iter_rows(min_row=2) if r[0].value}
    pairs, rows = [], []
    rescored = rescored_verdicts()
    for r in wb["check"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or r[0] not in key:
            continue
        k = key[r[0]]
        fact, group, unl = rescored.get(r[0], (k["scorer"]["fact"], k["group"], False))
        equity = fact == "correct, column not assessed"
        h = human_class(r[11], header_answer(r[0], r[12], r[10]), r[13], equity, unl)
        if h is None:
            continue
        s = "complete" if equity else group                     # row-only comparison for unassessed equity columns
        pairs.append((s, h)); rows.append(dict(item=r[0], system=k["system"], scorer=s, fact=fact, human=h, notes=r[14]))
    n = len(pairs)
    conf = Counter(pairs)
    out = ["# Human check of the fact scorer", "",
           f"Workbook: `{path.name}`. Checker: {meta.get('checker') or '(not given)'}; minutes: {meta.get('minutes spent') or '?'}; "
           f"confirmation of a check by hand: {meta.get('answer') or '(not given)'}.", "",
           f"Items answered: {n} of {len(key)}. Agreement: {sum(a == b for a, b in pairs)} of {n} "
           f"({100 * sum(a == b for a, b in pairs) / n:.1f}%); Cohen's kappa {kappa(pairs):.3f}.", "",
           "| Scorer \\ Human | " + " | ".join(CLASSES) + " | n |", "|---|" + "---:|" * (len(CLASSES) + 1)]
    for s in CLASSES[:3]:
        tot = sum(conf[(s, h)] for h in CLASSES)
        out.append(f"| {s} | " + " | ".join(str(conf[(s, h)]) for h in CLASSES) + f" | {tot} |")
    out += ["", "Scorer errors within each verdict group (exact 95% interval):", ""]
    for s in CLASSES[:3]:
        tot = sum(conf[(s, h)] for h in CLASSES)
        if tot:
            k = tot - conf[(s, s)]
            lo, hi = cp(k, tot)
            out.append(f"- {s}: {k} of {tot} disagree ({100 * k / tot:.1f}%, {100 * lo:.1f}–{100 * hi:.1f}%).")
    out += ["", "By system (agreement):", ""]
    for sysname in sorted({r["system"] for r in rows}):
        rs = [r for r in rows if r["system"] == sysname]
        out.append(f"- {sysname}: {sum(r['scorer'] == r['human'] for r in rs)} of {len(rs)}.")
    out += ["", "Disagreements:", "", "| Item | System | Scorer class | Human | Notes |", "|---|---|---|---|---|"]
    for r in rows:
        if r["scorer"] != r["human"]:
            out.append(f"| {r['item']} | {r['system']} | {r['fact']} | {r['human']} | {r['notes'] or ''} |")
    res = PAPER / "scorer_check" / "results"
    res.mkdir(parents=True, exist_ok=True)
    (res / "SCORER_CHECK.md").write_text("\n".join(out) + "\n", encoding="utf-8")
    print("\n".join(out))


if __name__ == "__main__":
    main()
