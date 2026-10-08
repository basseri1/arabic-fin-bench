"""Build the blind check of the ground truth's context (line items and periods): gt_ctx/packet and gt_ctx/packet.zip.

The independent re-transcription (gt_iaa) gave the annotator our rows, labels and periods and asked for the figures,
so it could not test the context. This packet does the reverse: each workbook keeps the row kinds, note references
and figures of the ground truth, and leaves empty every row label, every column header and every period. The
annotator reconstructs them from the page images alone. tools/score_context_check.py compares them with gt/.

Three filings are drawn at random, one per page format, from the filings not used by any earlier check (the
gt_audit and gt_iaa filings are excluded, since their labels have circulated). The seed is fixed before the draw.

usage: python tools/make_context_packet.py [--seed "gt_ctx 2026-10-02"] [--out gt_ctx]
"""
import argparse
import csv
import json
import random
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_workbook  # noqa: E402

PAPER = Path(__file__).resolve().parent.parent
EXCLUDE = {"almarai_FY2024", "bishah_FY2016", "halwani_FY2024",                  # gt_audit
           "alif_meem_yaa_FY2023", "sedco_capital_reit_FY2022", "al_khodari_FY2018"}  # gt_iaa
STRATA = ["text layer", "images in a digital PDF", "fully scanned"]
FILL = PatternFill("solid", fgColor="FFF2CC")                                      # cells to fill in
FIRST_ROW, VCOLS = 10, range(4, 18)                                              # rows from 10; values in D..Q

README_SHEET = [
    "Check of line items and periods — read the README.md of the packet first",
    "",
    "1. On the 'filing' sheet, write your name in 'annotator'.",
    "2. Each statement sheet already holds the figures of every printed line, in printed order, with the row kind",
    "   (section / item / subtotal / total) and the note references. The yellow cells are empty: that is your part.",
    "3. Column B (label_ar): for each row, find the printed line that carries these figures and type its label exactly",
    "   as printed, in Arabic. For a section row (no figures), type the heading printed at that point. A total printed",
    "   without a label stays empty.",
    "4. Row 7: type each value column's header as printed. Row 8: type the end of its period as YYYY-MM-DD (or the",
    "   year alone if no date is printed). In an equity statement, row 7 is the component heading (e.g. share capital).",
    "5. If the figures of a row are not on one printed line, or a column's figures belong to two periods, do not",
    "   guess: write FIGURES: or PERIOD: and an explanation in the comment column (T).",
    "6. Do not insert, delete or move rows or columns. Do not change the figures.",
]


def strata(inv):
    out = {s: [] for s in STRATA}
    for fid, r in inv.items():
        if fid in EXCLUDE:
            continue
        f = r["statement_pages_format"].lower()
        out["fully scanned" if "scanned" in f else "images in a digital PDF" if "images" in f else "text layer"].append(fid)
    return {s: sorted(v) for s, v in out.items()}


def blind(path, fid):
    """Empty labels, column headers and periods; drop sum_of and the live check; keep kinds, notes and figures."""
    wb = load_workbook(path)
    ws = wb["README"]
    ws.delete_rows(1, ws.max_row)
    for i, line in enumerate(README_SHEET, 1):
        ws.cell(row=i, column=1, value=line)
    fs = wb["filing"]
    for row in fs.iter_rows(min_row=2):
        if row[0].value in ("annotator", "verified_by", "comment"):
            row[1].value = None
        if row[0].value == "annotator":
            row[2].value = "your name"
    for ws in wb.worksheets[2:]:
        ws["B2"].value = None                                                    # statement title as printed
        for c in VCOLS:
            for r in (7, 8):
                ws.cell(row=r, column=c).value = None
                ws.cell(row=r, column=c).fill = FILL
        for r in range(FIRST_ROW, ws.max_row + 1):
            if ws.cell(row=r, column=1).value in (None, ""):
                continue
            ws.cell(row=r, column=2).value = None
            ws.cell(row=r, column=2).fill = FILL
            ws.cell(row=r, column=18).value = None                               # sum_of
            ws.cell(row=r, column=19).value = None                               # live check
            ws.cell(row=r, column=20).value = None                               # our comments hint at the line item
        ws.cell(row=7, column=18).value = None
        ws.cell(row=7, column=19).value = None
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", default="gt_ctx 2026-10-02")
    ap.add_argument("--out", default=str(PAPER / "gt_ctx"))
    a = ap.parse_args()
    out = Path(a.out)
    pk = out / "packet"
    if pk.exists():
        shutil.rmtree(pk)
    for d in ("workbooks", "pages", "filings"):
        (pk / d).mkdir(parents=True)
    (out / "returned").mkdir(exist_ok=True)
    (out / "results").mkdir(exist_ok=True)

    inv = {r["filing_id"]: r for r in csv.DictReader(open(PAPER / "dataset_inventory.csv", encoding="utf-8-sig"))}
    frame = strata(inv)
    rng = random.Random(a.seed)
    chosen = [(s, rng.choice(frame[s])) for s in STRATA]

    rows = []
    for stratum, fid in chosen:
        gt = json.load(open(PAPER / "gt" / f"{fid}.json", encoding="utf-8"))
        wbp = pk / "workbooks" / f"{fid}_context.xlsx"
        make_workbook.build(fid, draft=gt, out=str(wbp))
        blind(wbp, fid)
        pages = sorted({int(p) for st in gt["statements"] for p in st["pages"]})
        (pk / "pages" / fid).mkdir()
        for p in pages:
            subprocess.run(["pdftoppm", "-jpeg", "-r", "200", "-f", str(p), "-l", str(p), "-singlefile",
                            str(PAPER / "filings" / f"{fid}.pdf"), str(pk / "pages" / fid / f"page_{p:03d}")], check=True)
        shutil.copy(PAPER / "filings" / f"{fid}.pdf", pk / "filings" / f"{fid}.pdf")
        n_rows = sum(1 for st in gt["statements"] for r in st["rows"] if r.get("values"))
        n_cols = sum(len(st["columns"]) for st in gt["statements"])
        rows.append((stratum, fid, inv[fid]["proposed_split"].split()[0], len(gt["statements"]), n_rows, n_cols,
                     ", ".join(map(str, pages))))

    with open(pk / "time_log.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["filing_id", "date", "start", "end", "minutes", "notes"])
        for r in rows:
            w.writerow([r[1], "", "", "", "", ""])
    tbl = ["| Filing | Page format | Statements | Rows with figures | Value columns | PDF pages |", "|---|---|---:|---:|---:|---|"]
    tbl += [f"| {r[1]} | {r[0]} | {r[3]} | {r[4]} | {r[5]} | {r[6]} |" for r in rows]
    tbl.append(f"| **Total** | | **{sum(r[3] for r in rows)}** | **{sum(r[4] for r in rows)}** | **{sum(r[5] for r in rows)}** | |")
    (pk / "README.md").write_text("\n".join([
        "# Check of line items and periods: instructions", "",
        "Thank you for checking the benchmark's ground truth. This time the figures are given and the context is not: for "
        "every line of three Saudi financial filings you write the label printed on that line, and for every column its "
        "header and period. We compare your labels and periods with ours to measure how often the ground truth puts a "
        "figure under the wrong line item or period.", "",
        "## Rules that keep the check valid", "",
        "- **Work only from the page images or the PDFs in this folder.** Do not look at the project's ground-truth files, "
        "drafts, model outputs or earlier workbooks, and do not discuss the filings with the authors until you have finished.",
        "- **Do not use OCR or AI tools.** Read and type the labels yourself. Do not copy text out of the PDFs either: some "
        "text layers are garbled, and the check must be your own reading. This is reported as an independent annotation.",
        "- **Record it truthfully.** Write your name in `annotator` on each workbook's `filing` sheet, and your real start "
        "and end times in `time_log.csv`.",
        "- **Type what is printed, even when it looks wrong**, including spelling mistakes on the page.", "",
        "## What is in this folder", "",
        "- `workbooks/`: one Excel workbook per filing, one sheet per statement. Each row already holds the figures of one "
        "printed line, in printed order, with its row kind and note references. **Column B (labels) and rows 7–8 (column "
        "headers and periods) are empty: that is your part.**",
        "- `pages/`: the statement pages as images, named by their page number in the PDF (cell B3 of each sheet).",
        "- `filings/`: the original PDFs, if you prefer to zoom in a PDF viewer.",
        "- `time_log.csv`: note when you start and finish each filing.", "",
        "## The three filings", ""] + tbl + ["",
        "## How to fill a statement sheet", "",
        "1. Open the page shown in cell B3.",
        "2. Row 7: for each value column (D onwards), type the column header as printed, e.g. `31 ديسمبر 2023م`. Row 8: "
        "the end of that column's period as YYYY-MM-DD, e.g. `2023-12-31` (the year alone if the page gives no date). In a "
        "statement of changes in equity the columns are components of equity: type the heading in row 7 and leave row 8 "
        "empty unless a date is printed.",
        "3. Go down the rows. For each row, find the printed line whose figures match the row's figures and type its label "
        "in column B, exactly as printed (Arabic, including brackets and spelling variants). For a `section` row, type the "
        "heading printed at that point. A total printed under a rule without any label stays empty.",
        "4. If a row's figures are not on a single printed line, or you cannot find them, write `FIGURES:` and what you see "
        "in column T. If a column mixes two periods, write `PERIOD:` and an explanation in column T. Do not guess.",
        "5. Do not insert, delete or move rows or columns, and do not change any figure.", "",
        "Expect about 15 to 30 seconds per line, roughly two to three hours in all.", "",
        "When you have finished, return the three workbooks (keep their names) and `time_log.csv`."]) + "\n", encoding="utf-8")

    (out / "SAMPLING.md").write_text("\n".join([
        "# How the context-check sample was drawn", "",
        f"- **Frame:** the 32 filings except the six used by earlier checks of the ground truth ({', '.join(sorted(EXCLUDE))}), "
        "whose labels have already circulated.",
        "- **Strata:** the format of the statement pages, (a garbled text layer counts as text layer).",
        f"- **Draw:** one filing per stratum with `random.Random({a.seed!r}).choice` over the alphabetically sorted filing IDs, "
        "strata in the order text layer, images, scanned (`tools/make_context_packet.py`). The seed was fixed before the draw.", "",
        "| Stratum | Filings in frame | Chosen | Split | Statements | Rows with figures | Value columns |", "|---|---:|---|---|---:|---:|---:|"]
        + [f"| {r[0]} | {len(frame[r[0]])} | {r[1]} | {r[2]} | {r[3]} | {r[4]} | {r[5]} |" for r in rows]
        + ["", "Filings in each stratum:", ""] + [f"- **{s}:** {', '.join(v)}" for s, v in frame.items()]) + "\n", encoding="utf-8")

    zp = shutil.make_archive(str(out / "packet"), "zip", root_dir=out, base_dir="packet")
    print("chosen:", ", ".join(f"{fid} ({s})" for s, fid in chosen))
    print("wrote", pk, "and", zp)


if __name__ == "__main__":
    main()
