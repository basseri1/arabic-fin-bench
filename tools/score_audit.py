"""Score the blind ground-truth audit: compare the auditor's workbooks with gt/, cell by cell.

Every figure cell of the ground-truth structure is compared (the auditor's workbook has the same rows and columns, in
the same order). A disagreement is classed as: missing in audit, extra in audit, dash against number, sign, scale
(x10..x10,000), one digit (substituted, inserted, deleted or two neighbours swapped), unreadable entry, or other.
Rows the auditor added below the last row, values in columns the ground truth does not have, and tagged comments
(LABEL:, NOTE:, HEADER:, PERIOD:, PRINTED, MISSING ROW) are listed too.

Each disagreement goes to results/adjudication.csv for a decision against the page image: GT_WRONG, AUDIT_WRONG,
BOTH_WRONG or PRINTED (the page itself is inconsistent). Decisions already entered are kept when the script is
re-run. Once every value disagreement is decided, the report gives the ground-truth error rate among the audited
figures with its exact one-sided 95% upper bound (Clopper-Pearson), and a summary sentence.

usage: python tools/score_audit.py [--returned gt_audit/returned] [--results gt_audit/results]
"""
import argparse
import csv
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gtlib  # noqa: E402
import make_workbook as mw  # noqa: E402

PAPER = gtlib.PAPER
FIELDS = ["filing_id", "sheet", "statement", "pages", "excel_row", "row_label", "column", "period", "gt_value",
          "audit_value", "class", "auditor_comment", "decision", "correct_value", "adjudicator_note"]
DECISIONS = {"GT_WRONG", "AUDIT_WRONG", "BOTH_WRONG", "PRINTED"}
TAGS = ("LABEL:", "NOTE:", "HEADER:", "PERIOD:", "PRINTED", "MISSING ROW")
VALUE_CLASSES = {"missing in audit", "extra in audit", "dash against number", "sign", "scale", "one digit",
                 "unreadable entry", "other"}


def text(v):
    return "" if v is None else str(v).strip()


def is_zero(v):
    return v == "nil" or (isinstance(v, (int, float)) and v == 0)


def same(g, a):
    if g is None or a is None:
        return g is None and a is None
    if is_zero(g) or is_zero(a):
        return is_zero(g) and is_zero(a)
    return math.isclose(float(g), float(a), rel_tol=1e-12, abs_tol=1e-9)


def digits(v):
    return re.sub(r"[^0-9.]", "", f"{abs(v):f}".rstrip("0").rstrip(".") if isinstance(v, float) else str(abs(v)))


def one_edit(x, y):
    """True when y is x with one character substituted, inserted or deleted, or two neighbours swapped."""
    if x == y:
        return False
    if len(x) == len(y):
        diff = [i for i in range(len(x)) if x[i] != y[i]]
        return len(diff) == 1 or (len(diff) == 2 and diff[1] == diff[0] + 1 and x[diff[0]] == y[diff[1]] and x[diff[1]] == y[diff[0]])
    if abs(len(x) - len(y)) == 1:
        s, l = (x, y) if len(x) < len(y) else (y, x)
        return any(l[:i] + l[i + 1:] == s for i in range(len(l)))
    return False


def classify(g, a):
    if a is None:
        return "missing in audit"
    if g is None:
        return "extra in audit"
    if is_zero(g) != is_zero(a):
        return "dash against number"
    g, a = float(g), float(a)
    if math.isclose(abs(g), abs(a), rel_tol=1e-12, abs_tol=1e-9):
        return "sign"
    if g and a and (g > 0) == (a > 0):
        r = abs(a / g)
        for k in range(1, 5):
            if math.isclose(r, 10 ** k, rel_tol=1e-9) or math.isclose(r, 10 ** -k, rel_tol=1e-9):
                return "scale"
    if (g > 0) == (a > 0) and one_edit(digits(g), digits(a)):
        return "one digit"
    return "other"


def cp_upper(k, n, alpha=0.05):
    """Exact one-sided upper confidence bound for a binomial proportion (Clopper-Pearson)."""
    if n == 0:
        return float("nan")
    if k >= n:
        return 1.0
    cdf = lambda p: sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1))
    lo, hi = k / n, 1.0
    for _ in range(100):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if cdf(mid) > alpha else (lo, mid)
    return hi


def fmt(v):
    if v is None:
        return ""
    if v == "nil":
        return "-"
    return gtlib.fmt_value(v) if hasattr(gtlib, "fmt_value") else str(v)


def label(v):
    """A cell's category for agreement statistics: the value itself (dash = 0), or the raw entry if unreadable."""
    if v is None:
        return ""
    if is_zero(v):
        return "0"
    if isinstance(v, (int, float)):
        return f"{float(v):.6f}"
    return "?" + str(v)


def kappas(pairs):
    """Cohen's kappa for the two annotators, and Fleiss' kappa (which for two raters equals Scott's pi)."""
    n = len(pairs)
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe_c = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    pj = {k: (ca[k] + cb[k]) / (2 * n) for k in set(ca) | set(cb)}
    pe_f = sum(v * v for v in pj.values())
    return po, (po - pe_c) / (1 - pe_c), (po - pe_f) / (1 - pe_f)


def score_workbook(path):
    wb = load_workbook(path, data_only=False)
    meta = {text(r[0].value): r[1].value for r in wb["filing"].iter_rows(min_row=2) if r[0].value}
    fid = text(meta.get("filing_id"))
    gt = gtlib.load(PAPER / "gt" / f"{fid}.json")
    sheets = [ws for ws in wb.worksheets if ws.title not in ("README", "filing")]
    lines, counts, pairs = [], Counter(), []
    for ws, st in zip(sheets, gt["statements"]):
        cols = st["columns"]
        pages = ",".join(str(p) for p in st.get("pages", []))
        base = dict(filing_id=fid, sheet=ws.title, statement=st["type"], pages=pages)
        for j, row in enumerate(st["rows"]):
            x = mw.FIRST + j
            comment = text(ws[f"{mw.COMMENT_COL}{x}"].value)
            gv = row.get("values") or {}
            for k, cd in enumerate(cols[:len(mw.VCOLS)]):
                g = gv.get(cd["id"])
                raw = ws[f"{mw.VCOLS[k]}{x}"].value
                try:
                    a = gtlib.parse_value(raw)
                    cls = None if same(g, a) else classify(g, a)
                except ValueError:
                    a, cls = text(raw), "unreadable entry"
                if g is None and a is None:
                    continue
                counts["cells"] += 1
                pairs.append((label(g), label(a)))
                if isinstance(g, (int, float)):
                    counts["gt_figures"] += 1               # a figure, as counted by gtlib.stats
                elif g == "nil":
                    counts["gt_dashes"] += 1
                if cls is None:
                    counts["agree"] += 1
                    continue
                counts[cls] += 1
                lines.append(dict(base, excel_row=x, row_label=row.get("label_ar", ""), column=cd.get("label", ""),
                                  period=cd.get("period", ""), gt_value=fmt(g), audit_value=a if isinstance(a, str) else fmt(a),
                                  **{"class": cls}, auditor_comment=comment))
            if comment.upper().startswith(TAGS):
                lines.append(dict(base, excel_row=x, row_label=row.get("label_ar", ""), column="", period="",
                                  gt_value="", audit_value="", **{"class": "comment " + comment.split(":")[0].split()[0].upper()},
                                  auditor_comment=comment))
            for k in range(len(cols), len(mw.VCOLS)):           # values typed in a column the ground truth lacks
                if text(ws[f"{mw.VCOLS[k]}{x}"].value):
                    lines.append(dict(base, excel_row=x, row_label=row.get("label_ar", ""), column=text(ws[f"{mw.VCOLS[k]}7"].value),
                                      period="", gt_value="", audit_value=text(ws[f"{mw.VCOLS[k]}{x}"].value),
                                      **{"class": "extra column"}, auditor_comment=comment))
        for x in range(mw.FIRST + len(st["rows"]), mw.LAST + 1):   # rows the auditor added below the last row
            cells = [text(ws[f"{c}{x}"].value) for c in ["A", "B", "C"] + mw.VCOLS]
            if any(cells):
                lines.append(dict(base, excel_row=x, row_label=cells[1], column="", period="", gt_value="",
                                  audit_value=" | ".join(c for c in cells[3:] if c), **{"class": "added row"},
                                  auditor_comment=text(ws[f"{mw.COMMENT_COL}{x}"].value)))
    return fid, text(meta.get("annotator")), counts, lines, pairs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--returned", default=str(PAPER / "gt_audit" / "returned"))
    ap.add_argument("--results", default=str(PAPER / "gt_audit" / "results"))
    a = ap.parse_args()
    returned, results = Path(a.returned), Path(a.results)
    results.mkdir(parents=True, exist_ok=True)
    books = sorted(p for p in returned.glob("*.xlsx") if not p.name.startswith("~$"))
    if not books:
        raise SystemExit(f"no workbooks in {returned}")

    per, all_lines, auditors, all_pairs = {}, [], set(), []
    for p in books:
        fid, who, counts, lines, pairs = score_workbook(p)
        all_pairs += pairs
        per[fid] = counts
        all_lines += lines
        if who:
            auditors.add(who)

    # keep decisions already entered
    adj = results / "adjudication.csv"
    key = lambda d: (d["filing_id"], d["sheet"], str(d["excel_row"]), d["column"], d["class"])
    old = {}
    if adj.exists():
        for d in csv.DictReader(open(adj, encoding="utf-8-sig")):
            old[key(d)] = d
    for d in all_lines:
        prev = old.get(key(d))
        for f in ("decision", "correct_value", "adjudicator_note"):
            d[f] = prev.get(f, "") if prev else ""
    with open(adj, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for d in all_lines:
            w.writerow({k: d.get(k, "") for k in FIELDS})

    # time log, if returned
    minutes = {}
    tl = returned / "time_log.csv"
    if tl.exists():
        for r in csv.DictReader(open(tl, encoding="utf-8-sig")):
            try:
                minutes[r["filing_id"]] = minutes.get(r["filing_id"], 0) + float(r.get("minutes") or 0)
            except ValueError:
                pass

    tot = sum(per.values(), Counter())
    value_lines = [d for d in all_lines if d["class"] in VALUE_CLASSES]
    decided = [d for d in value_lines if d["decision"].strip().upper() in DECISIONS]
    dec = Counter(d["decision"].strip().upper() for d in decided)
    wrong = lambda d: d["decision"].strip().upper() in ("GT_WRONG", "BOTH_WRONG")
    fig_err = sum(1 for d in decided if wrong(d) and d["gt_value"] not in ("", "-"))     # a wrong figure in the GT
    omitted = sum(1 for d in decided if wrong(d) and d["gt_value"] == "")              # a printed figure the GT lacks
    dash_err = sum(1 for d in decided if wrong(d) and d["gt_value"] == "-")
    gt_err = fig_err + omitted
    pending = len(value_lines) - len(decided)
    other_lines = [d for d in all_lines if d["class"] not in VALUE_CLASSES]

    out = ["# Ground-truth audit: results", "",
           f"Workbooks scored: {len(books)} ({', '.join(sorted(per))}). Auditor: {', '.join(sorted(auditors)) or '(name not given)'}.", "",
           "## Agreement", "",
           "Cells compared: every figure cell where the ground truth or the auditor has a value.", "",
           "| Filing | Cells | Agree | Agreement | Disagreements | Minutes logged |", "|---|---:|---:|---:|---:|---:|"]
    for fid, c in sorted(per.items()):
        n, ag = c["cells"], c["agree"]
        out.append(f"| {fid} | {n} | {ag} | {100 * ag / n:.2f}% | {n - ag} | {minutes.get(fid, 0):.0f} |")
    out.append(f"| **Total** | **{tot['cells']}** | **{tot['agree']}** | **{100 * tot['agree'] / tot['cells']:.2f}%** | "
               f"**{tot['cells'] - tot['agree']}** | {sum(minutes.values()):.0f} |")
    po, kc, kf = kappas(all_pairs)
    out += ["", f"Inter-annotator agreement over the {len(all_pairs)} compared cells: raw agreement {100 * po:.2f}%, "
            f"Cohen's kappa {kc:.4f}, Fleiss' kappa {kf:.4f} (two annotators, so the pairwise Cohen's kappa is a single value; "
            "categories are the cell values, and exact figures rarely agree by chance, so both kappas are close to raw agreement)."]
    out += ["", "Disagreements by kind: " + (", ".join(f"{k} {tot[k]}" for k in sorted(VALUE_CLASSES) if tot[k]) or "none") + ".",
            "", f"Other items for review (added rows, extra columns, tagged comments): {len(other_lines)}"
            + (" (" + ", ".join(f"{k} {v}" for k, v in sorted(Counter(d['class'] for d in other_lines).items())) + ")." if other_lines else "."),
            "", "## Adjudication", ""]
    if pending:
        out.append(f"{pending} of {len(value_lines)} value disagreements still need a decision in `adjudication.csv` "
                   "(GT_WRONG, AUDIT_WRONG, BOTH_WRONG or PRINTED). Re-run this script after deciding.")
    else:
        n = tot["gt_figures"] + omitted
        ub = cp_upper(gt_err, n)
        out += [f"All {len(value_lines)} value disagreements decided: " + ", ".join(f"{k} {dec[k]}" for k in sorted(DECISIONS)) + ".", "",
                f"Ground-truth errors among the {n} audited figures: **{gt_err}** ({fig_err} wrong, {omitted} missing; "
                f"{100 * gt_err / n:.2f}%, exact one-sided 95% upper bound {100 * ub:.2f}%). Printed dashes audited: "
                f"{tot['gt_dashes']}, of which wrong in the ground truth: {dash_err}.", "",
                "**Summary:** A second annotator, blind to the ground truth, re-transcribed the "
                f"{n} figures of three filings drawn at random, one per page format; the two transcriptions agreed on "
                f"{100 * tot['agree'] / tot['cells']:.1f}\\% of cells, and adjudication against the page images found {gt_err} "
                f"ground-truth error{'s' if gt_err != 1 else ''} (one-sided 95\\% upper bound {100 * ub:.2f}\\% of figures)."]
    (results / "AUDIT_REPORT.md").write_text("\n".join(out) + "\n", encoding="utf-8")
    print("\n".join(out))
    print(f"\nwrote {results / 'AUDIT_REPORT.md'} and {adj}")


if __name__ == "__main__":
    main()
